from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher
from io import BytesIO
import math
import re
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as _Document
from docx.enum.text import WD_COLOR_INDEX
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


# Whole numerical expressions are matched before words. The expression covers
# dates, percentages, clause references, ranges, ratios, fractions, document
# numbers and common formula-like chains. A separated currency word/code remains
# a separate word, while an attached currency symbol stays with the number.
_CURRENCY_SYMBOLS = r"$€£¥₴₽₸₾₹₪"
_NUMERIC_RE = rf"""
    (?:\(\s*)?
    [{re.escape(_CURRENCY_SYMBOLS)}]?
    (?:№[ \t\u00A0\u202F]*)?
    (?:
        \d{{1,3}}(?:[ \u00A0\u202F]\d{{3}})+(?:[.,]\d+)?
        |
        \d+
        (?:
            [ \t\u00A0\u202F]*[.,:/\\+\-*×=–—‑][ \t\u00A0\u202F]*\d+
        )*
    )
    (?:[ \t\u00A0\u202F]*[%‰])?
    (?:[ \t\u00A0\u202F]*[-‑–—/][ \t\u00A0\u202F]*[^\W_]+)?
    [{re.escape(_CURRENCY_SYMBOLS)}]?
    (?:\s*\))?
"""

# A word includes internal apostrophes and hyphens. This guarantees that a
# changed letter, ending, apostrophe or hyphenated element highlights the whole
# word rather than an internal character.
_WORD_RE = r"[^\W_]+(?:[’'ʼ\-‑–—][^\W_]+)*"

_TOKEN_RE = re.compile(
    rf"(?P<number>{_NUMERIC_RE})|(?P<word>{_WORD_RE})|(?P<punct>[^\w\s])",
    re.UNICODE | re.VERBOSE,
)

_SERVICE_LABELS_CASEFOLDED = {
    "Редакція Покупця".casefold(),
    "Редакція Постачальника".casefold(),
    "Редакція Замовника".casefold(),
    "Редакція Виконавця".casefold(),
}

_SERVICE_LABEL_FINAL_PUNCTUATION_RE = re.compile(
    r"\s*[.:;,\-‐‑‒–—−]$",
    re.UNICODE,
)

_NUMBERING_HEADER_RE = re.compile(
    r"""
    ^
    (?:
        №
        (?:\s*
            (?:
                ст(?:\.|атті|атьи)?
                |
                п(?:\.|ункт(?:у|а)?)?
                |
                п/?п
                |
                з/?п
                |
                item
                |
                clause
            )
        )?
        (?:\s*(?:договору|договора|contract))?
        \.?
        |
        номер(?:\s+(?:статті|статьи|пункту|пункта|договору|договора))?
        |
        п/?п
        |
        з/?п
        |
        пункт
        |
        п\.?
        |
        item
        |
        clause
        |
        no\.?
        |
        n
    )
    $
    """,
    re.IGNORECASE | re.VERBOSE,
)

_NUMBERING_VALUE_RE = re.compile(
    r"^(?:№\s*)?\d+(?:[.\-]\d+)*(?:[.)])?$",
    re.UNICODE,
)

_HEADER_KEYWORDS = (
    "редакц",
    "базов",
    "постачальник",
    "покупець",
    "покупця",
    "supplier",
    "buyer",
    "base version",
    "wording",
    "редакция",
)


@dataclass(frozen=True)
class Token:
    text: str
    start: int
    end: int
    kind: str


@dataclass
class ProcessingStats:
    tables_found: int = 0
    rows_checked: int = 0
    rows_skipped: int = 0
    header_rows_skipped: int = 0
    numbering_columns_excluded: int = 0
    third_labels_excluded: int = 0
    cells_changed: int = 0
    highlighted_fragments: int = 0

    def as_dict(self) -> dict[str, int]:
        return {
            "tables_found": self.tables_found,
            "rows_checked": self.rows_checked,
            "rows_skipped": self.rows_skipped,
            "header_rows_skipped": self.header_rows_skipped,
            "numbering_columns_excluded": self.numbering_columns_excluded,
            "third_labels_excluded": self.third_labels_excluded,
            "cells_changed": self.cells_changed,
            "highlighted_fragments": self.highlighted_fragments,
        }


def tokenize(text: str) -> list[Token]:
    """Tokenize into complete numbers, complete words and separate punctuation.

    Whitespace is intentionally omitted from comparison and can therefore never
    be highlighted, while its original placement is preserved in the document.
    """
    return [
        Token(
            text=match.group(0),
            start=match.start(),
            end=match.end(),
            kind=match.lastgroup or "punct",
        )
        for match in _TOKEN_RE.finditer(text)
    ]


def merge_spans(spans: Iterable[tuple[int, int]]) -> list[tuple[int, int]]:
    normalized = sorted((start, end) for start, end in spans if end > start)
    if not normalized:
        return []

    result = [normalized[0]]
    for start, end in normalized[1:]:
        previous_start, previous_end = result[-1]
        # Merge only overlapping or immediately adjacent units. Spaces remain
        # outside highlighting because they create a positive gap.
        if start <= previous_end:
            result[-1] = (previous_start, max(previous_end, end))
        else:
            result.append((start, end))
    return result


def _relative_position(index: int, length: int) -> float:
    if length <= 1:
        return 0.0
    return index / (length - 1)


def _block_is_contextual(
    block,
    all_blocks,
    left_tokens: Sequence[Token],
    right_tokens: Sequence[Token],
) -> bool:
    """Reject isolated accidental matches that lack order/context support."""
    if block.size >= 2:
        return True

    i, j = block.a, block.b
    left_len = len(left_tokens)
    right_len = len(right_tokens)
    token = left_tokens[i]

    if left_len == right_len == 1:
        return True

    at_same_start = i == 0 and j == 0
    at_same_end = i == left_len - 1 and j == right_len - 1
    if at_same_start or at_same_end:
        return True

    position_gap = abs(
        _relative_position(i, left_len) - _relative_position(j, right_len)
    )

    # A nearby multi-token anchor supplies enough corresponding context.
    anchors = [other for other in all_blocks if other.size >= 2]
    for other in anchors:
        left_gap = min(abs(i - (other.a - 1)), abs(i - (other.a + other.size)))
        right_gap = min(abs(j - (other.b - 1)), abs(j - (other.b + other.size)))
        if left_gap <= 2 and right_gap <= 2:
            return True

    # A singleton between the same preceding and following anchors is still
    # contextual even when one side contains a short inserted phrase. This
    # preserves shared bridge words such as "за" in:
    # "... ціни не менш, ніж за 14 днів" / "... ціни за 3 дні".
    previous = [
        other for other in anchors
        if other.a + other.size <= i and other.b + other.size <= j
    ]
    following = [
        other for other in anchors
        if other.a > i and other.b > j
    ]
    if previous and following:
        prev = max(previous, key=lambda other: (other.a + other.size, other.b + other.size))
        nxt = min(following, key=lambda other: (other.a, other.b))
        gaps = (
            i - (prev.a + prev.size),
            j - (prev.b + prev.size),
            nxt.a - (i + 1),
            nxt.b - (j + 1),
        )
        if max(gaps) <= 6:
            return True

    if token.kind == "punct":
        # Sentence-ending punctuation can correspond even when one version
        # adds or removes a trailing phrase/sentence.
        if i == left_len - 1 or j == right_len - 1:
            for other in anchors:
                left_gap = i - (other.a + other.size)
                right_gap = j - (other.b + other.size)
                if 0 <= left_gap <= 6 and 0 <= right_gap <= 6:
                    return True
        return position_gap <= 0.04 and abs(i - j) <= 1

    # Longer words/numbers can stand alone only when their positions correspond
    # closely. Short common tokens need even tighter positional agreement.
    threshold = 0.08 if len(token.text) >= 5 else 0.035
    return position_gap <= threshold and abs(i - j) <= 1


def _contextual_matching_blocks(
    left_tokens: Sequence[Token],
    right_tokens: Sequence[Token],
):
    matcher = SequenceMatcher(
        None,
        [token.text for token in left_tokens],
        [token.text for token in right_tokens],
        autojunk=False,
    )
    candidates = [block for block in matcher.get_matching_blocks() if block.size]
    return [
        block
        for block in candidates
        if _block_is_contextual(block, candidates, left_tokens, right_tokens)
    ]


def differing_spans(
    left: str,
    right: str,
) -> tuple[list[tuple[int, int]], list[tuple[int, int]]]:
    """Return character spans for exact differing complete units on both sides."""
    left_tokens = tokenize(left)
    right_tokens = tokenize(right)

    left_matched = [False] * len(left_tokens)
    right_matched = [False] * len(right_tokens)

    for block in _contextual_matching_blocks(left_tokens, right_tokens):
        for offset in range(block.size):
            left_matched[block.a + offset] = True
            right_matched[block.b + offset] = True

    def visible_spans(token: Token) -> list[tuple[int, int]]:
        # Even inside a complete numerical unit, whitespace is preserved but
        # never receives highlighting.
        return [
            (token.start + match.start(), token.start + match.end())
            for match in re.finditer(r"\S+", token.text, re.UNICODE)
        ]

    left_spans = [
        span
        for token, matched in zip(left_tokens, left_matched)
        if not matched
        for span in visible_spans(token)
    ]
    right_spans = [
        span
        for token, matched in zip(right_tokens, right_matched)
        if not matched
        for span in visible_spans(token)
    ]

    return merge_spans(left_spans), merge_spans(right_spans)


def cell_text(cell: _Cell) -> str:
    return "\n".join(paragraph.text for paragraph in cell.paragraphs)


def _iter_tables(parent: _Document | _Cell) -> Iterable[Table]:
    for table in parent.tables:
        yield table
        for row in table.rows:
            seen_cells: set[int] = set()
            for cell in row.cells:
                identity = id(cell._tc)
                if identity in seen_cells:
                    continue
                seen_cells.add(identity)
                yield from _iter_tables(cell)


def _normalize_visible_whitespace(text: str) -> str:
    return re.sub(r"[\s\u00A0\u202F]+", " ", text).strip()


def is_excluded_third_label(text: str) -> bool:
    """Check whether the third-version cell contains only a service label."""
    normalized = _normalize_visible_whitespace(text)
    normalized = _SERVICE_LABEL_FINAL_PUNCTUATION_RE.sub("", normalized, count=1)
    normalized = normalized.rstrip()
    return normalized.casefold() in _SERVICE_LABELS_CASEFOLDED


def _looks_like_numbering_column(table: Table) -> bool:
    if not table.rows or max(len(row.cells) for row in table.rows) < 4:
        return False

    values: list[str] = []
    for row in table.rows[:100]:
        if not row.cells:
            continue
        value = _normalize_visible_whitespace(cell_text(row.cells[0]))
        if value:
            values.append(value)

    if not values:
        return False

    if _NUMBERING_HEADER_RE.fullmatch(values[0]):
        return True

    number_like = sum(bool(_NUMBERING_VALUE_RE.fullmatch(value)) for value in values)
    if len(values) == 1:
        return number_like == 1

    required = max(2, math.ceil(len(values) * 0.60))
    return number_like >= required


def version_column_indices(table: Table) -> tuple[int, int, int] | None:
    if not table.rows:
        return None

    max_columns = max(len(row.cells) for row in table.rows)
    if max_columns < 3:
        return None

    if _looks_like_numbering_column(table):
        if max_columns < 4:
            return None
        return (1, 2, 3)

    return (0, 1, 2)


def _looks_like_header_row(
    texts: Sequence[str],
    *,
    row_index: int,
) -> bool:
    if row_index > 1:
        return False

    normalized = [_normalize_visible_whitespace(text).lower() for text in texts]
    nonempty = [text for text in normalized if text]
    if not nonempty:
        return True

    if any(len(text) > 100 for text in nonempty):
        return False

    keyword_cells = sum(
        any(keyword in text for keyword in _HEADER_KEYWORDS)
        for text in nonempty
    )
    return keyword_cells >= 2


def _paragraph_offsets(cell: _Cell) -> list[tuple[Paragraph, int, int]]:
    offsets: list[tuple[Paragraph, int, int]] = []
    cursor = 0
    for index, paragraph in enumerate(cell.paragraphs):
        text = paragraph.text
        offsets.append((paragraph, cursor, cursor + len(text)))
        cursor += len(text)
        if index < len(cell.paragraphs) - 1:
            cursor += 1  # virtual newline used by cell_text()
    return offsets


def _overlapping_ranges(
    spans: Sequence[tuple[int, int]],
    paragraph_start: int,
    paragraph_end: int,
) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    for start, end in spans:
        overlap_start = max(start, paragraph_start)
        overlap_end = min(end, paragraph_end)
        if overlap_end > overlap_start:
            ranges.append(
                (overlap_start - paragraph_start, overlap_end - paragraph_start)
            )
    return merge_spans(ranges)


def _segment_is_highlighted(
    start: int,
    end: int,
    highlight_ranges: Sequence[tuple[int, int]],
) -> bool:
    return any(
        start < range_end and end > range_start
        for range_start, range_end in highlight_ranges
    )


def _clone_run_with_text(run: Run, paragraph: Paragraph, text: str) -> Run:
    new_r = deepcopy(run._r)

    # Retain all run formatting. Rebuild textual content only for the cloned
    # segment so the visible styling remains unchanged.
    for child in list(new_r):
        if child is not new_r.rPr:
            new_r.remove(child)

    new_run = Run(new_r, paragraph)
    new_run.text = text
    return new_run


def _highlight_paragraph(
    paragraph: Paragraph,
    ranges: Sequence[tuple[int, int]],
) -> int:
    if not ranges:
        return 0

    ranges = merge_spans(ranges)
    run_positions: list[tuple[Run, int, int]] = []
    cursor = 0

    for run in paragraph.runs:
        run_text = run.text
        start = cursor
        end = start + len(run_text)
        run_positions.append((run, start, end))
        cursor = end

    fragments_added = 0

    for run, run_start, run_end in run_positions:
        if run_end <= run_start:
            continue

        relevant = merge_spans(
            (
                max(start, run_start),
                min(end, run_end),
            )
            for start, end in ranges
            if end > run_start and start < run_end
        )
        if not relevant:
            continue

        if (
            len(relevant) == 1
            and relevant[0][0] <= run_start
            and relevant[0][1] >= run_end
        ):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            fragments_added += 1
            continue

        boundaries = {run_start, run_end}
        for start, end in relevant:
            boundaries.add(start)
            boundaries.add(end)
        ordered = sorted(boundaries)

        parent = run._r.getparent()
        insertion_index = parent.index(run._r)
        original_text = run.text

        for absolute_start, absolute_end in zip(ordered, ordered[1:]):
            if absolute_end <= absolute_start:
                continue

            local_start = absolute_start - run_start
            local_end = absolute_end - run_start
            segment_text = original_text[local_start:local_end]
            if not segment_text:
                continue

            new_run = _clone_run_with_text(run, paragraph, segment_text)
            if _segment_is_highlighted(
                absolute_start,
                absolute_end,
                relevant,
            ):
                new_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                fragments_added += 1

            parent.insert(insertion_index, new_run._r)
            insertion_index += 1

        parent.remove(run._r)

    return fragments_added


def apply_highlights(cell: _Cell, spans: Sequence[tuple[int, int]]) -> int:
    fragments = 0
    for paragraph, paragraph_start, paragraph_end in _paragraph_offsets(cell):
        local_ranges = _overlapping_ranges(
            spans,
            paragraph_start,
            paragraph_end,
        )
        fragments += _highlight_paragraph(paragraph, local_ranges)
    return fragments


def compare_version_cells(
    cells: Sequence[_Cell],
) -> tuple[list[list[tuple[int, int]]], bool]:
    """Compare base↔second and, unless excluded, base↔third."""
    if len(cells) < 3:
        raise ValueError("At least three version cells are required.")

    texts = [cell_text(cell) for cell in cells[:3]]
    all_spans: list[list[tuple[int, int]]] = [[], [], []]

    base_second, second = differing_spans(texts[0], texts[1])
    all_spans[0].extend(base_second)
    all_spans[1].extend(second)

    third_excluded = is_excluded_third_label(texts[2])
    if not third_excluded:
        base_third, third = differing_spans(texts[0], texts[2])
        all_spans[0].extend(base_third)
        all_spans[2].extend(third)

    return [merge_spans(spans) for spans in all_spans], third_excluded


def _document_text_snapshot(document: _Document) -> tuple:
    snapshot = []
    for table in _iter_tables(document):
        table_rows = []
        for row in table.rows:
            table_rows.append(tuple(cell_text(cell) for cell in row.cells))
        snapshot.append(tuple(table_rows))
    return tuple(snapshot)


def process_docx(source: bytes) -> tuple[bytes, ProcessingStats]:
    """Add yellow highlighting to an editable DOCX without changing its text."""
    if not source:
        raise ValueError("Загруженный файл пуст.")

    try:
        document = Document(BytesIO(source))
    except Exception as exc:
        raise ValueError(
            "Файл не удалось открыть как редактируемый документ DOCX."
        ) from exc

    original_snapshot = _document_text_snapshot(document)
    stats = ProcessingStats()
    tables = list(_iter_tables(document))
    stats.tables_found = len(tables)

    for table in tables:
        indices = version_column_indices(table)
        if indices is None:
            stats.rows_skipped += len(table.rows)
            continue

        if indices[0] == 1:
            stats.numbering_columns_excluded += 1

        for row_index, row in enumerate(table.rows):
            if len(row.cells) <= max(indices):
                stats.rows_skipped += 1
                continue

            cells = [row.cells[index] for index in indices]

            # Horizontal merges can expose the same underlying cell at multiple
            # logical indexes. Such a row cannot be compared reliably.
            if len({id(cell._tc) for cell in cells}) < 3:
                stats.rows_skipped += 1
                continue

            texts = [cell_text(cell) for cell in cells]
            if _looks_like_header_row(texts, row_index=row_index):
                stats.header_rows_skipped += 1
                continue

            stats.rows_checked += 1
            cell_spans, third_excluded = compare_version_cells(cells)
            if third_excluded:
                stats.third_labels_excluded += 1

            for cell, spans in zip(cells, cell_spans):
                if not spans:
                    continue
                stats.cells_changed += 1
                stats.highlighted_fragments += apply_highlights(cell, spans)

    if _document_text_snapshot(document) != original_snapshot:
        raise RuntimeError(
            "Проверка целостности не пройдена: текст документа изменился."
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue(), stats
