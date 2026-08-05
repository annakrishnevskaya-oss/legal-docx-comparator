from io import BytesIO

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from comparator import (
    differing_spans,
    is_excluded_third_label,
    process_docx,
    tokenize,
)


def make_docx(rows, columns=3):
    document = Document()
    table = document.add_table(rows=len(rows), cols=columns)

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def highlighted_text(cell):
    parts = []
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                parts.append(run.text)
    return "".join(parts)


def all_text(document):
    return [
        [[paragraph.text for paragraph in cell.paragraphs] for cell in row.cells]
        for table in document.tables
        for row in table.rows
    ]


def test_identical_text_has_no_highlight():
    source = make_docx([("Оплата 5 днів", "Оплата 5 днів", "Оплата 5 днів")])
    result, stats = process_docx(source)

    document = Document(BytesIO(result))
    row = document.tables[0].rows[0]

    assert [highlighted_text(cell) for cell in row.cells] == ["", "", ""]
    assert stats.cells_changed == 0


def test_second_and_third_are_compared_only_with_base():
    source = make_docx([
        ("Оплата 5 днів", "Оплата 10 днів", "Оплата 5 днів"),
    ])
    result, _ = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [highlighted_text(cell) for cell in row.cells] == ["5", "10", ""]


def test_base_highlight_is_union_of_two_comparisons():
    source = make_docx([
        (
            "Оплата 5 робочих днів",
            "Оплата 10 робочих днів",
            "Оплата 5 календарних днів",
        ),
    ])
    result, _ = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    values = [highlighted_text(cell) for cell in row.cells]

    assert "5" in values[0]
    assert "робочих" in values[0]
    assert values[1] == "10"
    assert values[2] == "календарних"


def test_whole_words_are_highlighted():
    left, right = differing_spans("payment", "payments")
    assert left == [(0, 7)]
    assert right == [(0, 8)]

    left, right = differing_spans("Покупець", "Покупця")
    assert left == [(0, 8)]
    assert right == [(0, 7)]


def test_whole_numerical_expressions_are_highlighted():
    examples = [
        ("10", "15", "10", "15"),
        ("0.1%", "0.15%", "0.1%", "0.15%"),
        ("01.06.2026", "05.06.2026", "01.06.2026", "05.06.2026"),
        ("п. 5.2", "п. 5.3", "5.2", "5.3"),
        ("№ 123-А", "№ 124-А", "№123-А", "№124-А"),
        ("10 – 15", "10 – 20", "10–15", "10–20"),
    ]

    for left_text, right_text, expected_left, expected_right in examples:
        left, right = differing_spans(left_text, right_text)
        assert "".join(left_text[start:end] for start, end in left) == expected_left
        assert "".join(right_text[start:end] for start, end in right) == expected_right


def test_separated_currency_word_is_not_highlighted_when_unchanged():
    left, right = differing_spans("1 000,00 грн", "1 500,00 грн")
    assert "".join("1 000,00 грн"[start:end] for start, end in left) == "1000,00"
    assert "".join("1 500,00 грн"[start:end] for start, end in right) == "1500,00"


def test_attached_currency_symbol_is_part_of_number():
    left, right = differing_spans("₴100", "₴150")
    assert left == [(0, 4)]
    assert right == [(0, 4)]


def test_separate_punctuation_only():
    left, right = differing_spans("Текст.", "Текст,")
    assert left == [(5, 6)]
    assert right == [(5, 6)]


def test_whitespace_only_differences_are_ignored():
    left, right = differing_spans(
        "Оплата\t5 днів\nпісля поставки",
        "Оплата  5 днів після поставки",
    )
    assert left == []
    assert right == []


def test_missing_text_is_highlighted_only_where_present():
    source = make_docx([
        ("Оплата після поставки", "Оплата", "Оплата після поставки"),
    ])
    result, _ = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    values = [highlighted_text(cell) for cell in row.cells]

    assert values[0] == "післяпоставки"
    assert values[1] == ""
    assert values[2] == ""


def test_empty_cell_remains_empty():
    source = make_docx([
        ("Оплата після поставки.", "", "Оплата після поставки."),
    ])
    result, _ = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    assert highlighted_text(row.cells[0]) == "Оплатапісляпоставки."
    assert row.cells[1].text == ""
    assert highlighted_text(row.cells[1]) == ""
    assert highlighted_text(row.cells[2]) == ""


def test_service_label_is_case_insensitive_and_accepts_one_final_punctuation():
    accepted = [
        "Редакція Покупця",
        "редакція покупця",
        "РЕДАКЦІЯ ПОКУПЦЯ",
        " \tРедакція\u00A0Покупця.\n",
        "редакція постачальника:",
        "РЕДАКЦІЯ ПОСТАЧАЛЬНИКА;",
        "Редакція Покупця,",
        "Редакція Покупця -",
        "Редакція Покупця–",
        "Редакція Покупця —",
        "Редакція Замовника",
        "редакція замовника:",
        "РЕДАКЦІЯ ВИКОНАВЦЯ;",
        "Редакція Виконавця —",
    ]
    for value in accepted:
        assert is_excluded_third_label(value), value

    rejected = [
        "Редакція Покупця..",
        "Редакція Покупця:;",
        "Редакція Покупця: текст",
        "Редакція Покупця додаткова умова",
    ]
    for value in rejected:
        assert not is_excluded_third_label(value), value


def test_service_label_does_not_highlight_base_or_third():
    source = make_docx([
        (
            "Оплата 5 днів",
            "Оплата 10 днів",
            " \tРедакція\u00A0Покупця.\n",
        ),
    ])
    result, stats = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [highlighted_text(cell) for cell in row.cells] == ["5", "10", ""]
    assert row.cells[2].text == " \tРедакція\u00A0Покупця.\n"
    assert stats.third_labels_excluded == 1


def test_service_label_with_extra_contract_text_is_not_excluded():
    source = make_docx([
        (
            "Оплата 5 днів",
            "Оплата 5 днів",
            "Редакція Покупця: Оплата 5 днів",
        ),
    ])
    result, stats = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    assert highlighted_text(row.cells[2]) != ""
    assert stats.third_labels_excluded == 0


def test_numbering_column_is_excluded():
    source = make_docx(
        [
            ("№", "Базова редакція", "Друга редакція", "Третя редакція"),
            ("1", "Оплата 5 днів", "Оплата 10 днів", "Оплата 5 днів"),
            ("2.1", "Строк 3 дні", "Строк 3 дні", "Строк 4 дні"),
        ],
        columns=4,
    )
    result, stats = process_docx(source)

    table = Document(BytesIO(result)).tables[0]
    assert highlighted_text(table.rows[0].cells[0]) == ""
    assert [highlighted_text(cell) for cell in table.rows[1].cells] == [
        "",
        "5",
        "10",
        "",
    ]
    assert [highlighted_text(cell) for cell in table.rows[2].cells] == [
        "",
        "3",
        "",
        "4",
    ]
    assert stats.numbering_columns_excluded == 1


def test_header_row_is_not_highlighted():
    source = make_docx([
        ("Базова редакція", "Редакція Постачальника", "Редакція Покупця"),
        ("Оплата 5 днів", "Оплата 10 днів", "Оплата 5 днів"),
    ])
    result, stats = process_docx(source)

    table = Document(BytesIO(result)).tables[0]
    assert [highlighted_text(cell) for cell in table.rows[0].cells] == ["", "", ""]
    assert stats.header_rows_skipped == 1


def test_original_text_and_basic_run_formatting_are_preserved():
    document = Document()
    table = document.add_table(rows=1, cols=3)

    for index, text in enumerate(("Оплата 5 днів", "Оплата 10 днів", "Оплата 5 днів")):
        paragraph = table.cell(0, index).paragraphs[0]
        paragraph.clear()
        run = paragraph.add_run(text)
        run.bold = True
        run.italic = True

    before = BytesIO()
    document.save(before)

    result, _ = process_docx(before.getvalue())
    processed = Document(BytesIO(result))

    assert all_text(processed) == all_text(document)
    for cell in processed.tables[0].rows[0].cells:
        assert all(run.bold for run in cell.paragraphs[0].runs if run.text)
        assert all(run.italic for run in cell.paragraphs[0].runs if run.text)


def test_tokenizer_keeps_hyphenated_and_apostrophe_words_whole():
    assert [token.text for token in tokenize("слово-слово О’Браєн")] == [
        "слово-слово",
        "О’Браєн",
    ]


def test_case_variant_service_label_leaves_third_cell_completely_unchanged():
    original_label = " \tрЕдАкЦіЯ\u00A0пОкУпЦя —\n"
    source = make_docx([
        (
            "Оплата 5 днів",
            "Оплата 10 днів",
            original_label,
        ),
    ])
    result, stats = process_docx(source)

    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [highlighted_text(cell) for cell in row.cells] == ["5", "10", ""]
    assert row.cells[2].text == original_label
    assert stats.third_labels_excluded == 1


def test_new_service_labels_are_excluded_from_third_version():
    for label in (
        "Редакція Замовника",
        "редакція замовника:",
        "РЕДАКЦІЯ ВИКОНАВЦЯ;",
        "Редакція Виконавця —",
    ):
        source = make_docx([
            (
                "Строк 5 днів",
                "Строк 10 днів",
                label,
            ),
        ])
        result, stats = process_docx(source)

        row = Document(BytesIO(result)).tables[0].rows[0]
        assert [highlighted_text(cell) for cell in row.cells] == ["5", "10", ""]
        assert row.cells[2].text == label
        assert stats.third_labels_excluded == 1
